from flask import Flask, request, jsonify, send_file
from flask_cors import CORS
from dotenv import load_dotenv
from datetime import datetime
from apscheduler.schedulers.background import BackgroundScheduler
import os
import csv
import io
from openpyxl import Workbook

from database_enhanced import Database
from youtube_service import YouTubeService
from models import Video
from bulk_import import BulkImporter
from email_service import EmailService

# Load environment variables
load_dotenv()

app = Flask(__name__)
app.config['SECRET_KEY'] = os.getenv('FLASK_SECRET_KEY', 'dev-secret-key')
CORS(app)

# Initialize database
db = Database(os.getenv('DATABASE_PATH', 'videos.db'))

# Initialize YouTube service
try:
    youtube_service = YouTubeService(os.getenv('YOUTUBE_API_KEY', ''))
    API_CONFIGURED = True
except ValueError as e:
    youtube_service = None
    API_CONFIGURED = False
    print(f"WARNING: {str(e)}")

# Initialize email service
email_service = EmailService()

# Admin password
ADMIN_PASSWORD = os.getenv('ADMIN_PASSWORD', 'admin123')

# Debug: Print loaded password on startup
print("=" * 50)
print(f"ADMIN PASSWORD LOADED: '{ADMIN_PASSWORD}'")
print(f"EMAIL SERVICE: {'Enabled' if email_service.enabled else 'Disabled'}")
print("=" * 50)

# Scheduler for automatic searches
scheduler = BackgroundScheduler()
scheduler.start()

# Auth middleware
def check_auth(password):
    return password == ADMIN_PASSWORD

@app.route('/api/health', methods=['GET'])
def health_check():
    return jsonify({
        'status': 'ok',
        'api_configured': API_CONFIGURED,
        'database_ok': os.path.exists(db.db_path),
        'email_configured': email_service.enabled
    })

@app.route('/api/auth/login', methods=['POST'])
def login():
    data = request.json
    password = data.get('password', '')
    
    # Debug logging
    print(f"Login attempt - Received password: '{password}'")
    print(f"Expected password: '{ADMIN_PASSWORD}'")
    print(f"Passwords match: {password == ADMIN_PASSWORD}")
    
    if check_auth(password):
        return jsonify({'success': True, 'message': 'Authentication successful'})
    else:
        return jsonify({'success': False, 'message': 'Invalid password'}), 401

# ============================================================================
# ARTISTS ENDPOINTS
# ============================================================================

@app.route('/api/artists', methods=['GET'])
def get_artists():
    artists = db.get_all_artists()
    return jsonify([a.to_dict() for a in artists])

@app.route('/api/artists', methods=['POST'])
def add_artist():
    data = request.json
    name = data.get('name', '').strip()
    email = data.get('email', '').strip()
    contact_person = data.get('contact_person', '').strip()
    notes = data.get('notes', '').strip()
    
    if not name:
        return jsonify({'error': 'Artist name is required'}), 400
    
    artist_id = db.add_artist(name, email, contact_person, notes)
    
    if artist_id:
        return jsonify({'success': True, 'id': artist_id, 'message': 'Artist added'})
    else:
        return jsonify({'error': 'Artist already exists'}), 400

@app.route('/api/artists/<int:artist_id>', methods=['GET'])
def get_artist(artist_id):
    artist = db.get_artist(artist_id)
    if artist:
        return jsonify(artist.to_dict())
    else:
        return jsonify({'error': 'Artist not found'}), 404

@app.route('/api/artists/<int:artist_id>', methods=['PUT'])
def update_artist(artist_id):
    data = request.json
    success = db.update_artist(
        artist_id,
        name=data.get('name'),
        email=data.get('email'),
        contact_person=data.get('contact_person'),
        notes=data.get('notes'),
        active=data.get('active')
    )
    
    if success:
        return jsonify({'success': True, 'message': 'Artist updated'})
    else:
        return jsonify({'error': 'Artist not found'}), 404

@app.route('/api/artists/<int:artist_id>', methods=['DELETE'])
def delete_artist(artist_id):
    success = db.delete_artist(artist_id)
    
    if success:
        return jsonify({'success': True, 'message': 'Artist deleted'})
    else:
        return jsonify({'error': 'Artist not found'}), 404

# ============================================================================
# KEYWORDS ENDPOINTS (Enhanced)
# ============================================================================

@app.route('/api/keywords', methods=['GET'])
def get_keywords():
    artist_id = request.args.get('artist_id', type=int)
    keywords = db.get_all_keywords(artist_id)
    return jsonify([k.to_dict() for k in keywords])

@app.route('/api/keywords', methods=['POST'])
def add_keyword():
    data = request.json
    keyword = data.get('keyword', '').strip()
    artist_id = data.get('artist_id')
    auto_flag = data.get('auto_flag', False)
    priority = data.get('priority', 'Medium')
    
    if not keyword:
        return jsonify({'error': 'Keyword is required'}), 400
    
    keyword_id = db.add_keyword(keyword, artist_id, auto_flag, priority)
    
    if keyword_id:
        return jsonify({'success': True, 'id': keyword_id, 'message': 'Keyword added'})
    else:
        return jsonify({'error': 'Keyword already exists'}), 400

@app.route('/api/keywords/bulk-import', methods=['POST'])
def bulk_import_keywords():
    if 'file' not in request.files:
        return jsonify({'error': 'No file provided'}), 400
    
    file = request.files['file']
    if file.filename == '':
        return jsonify({'error': 'No file selected'}), 400
    
    try:
        # Read file content
        if file.filename.endswith('.csv'):
            content = file.read().decode('utf-8')
            keywords_data = BulkImporter.parse_csv(content)
        elif file.filename.endswith('.xlsx'):
            # Save temporarily
            temp_path = f"/tmp/{file.filename}"
            file.save(temp_path)
            keywords_data = BulkImporter.parse_excel(temp_path)
            os.remove(temp_path)
        else:
            return jsonify({'error': 'Invalid file format. Use CSV or XLSX'}), 400
        
        # Get artist mapping
        artists = db.get_all_artists()
        artist_map = {a.name: a.id for a in artists}
        
        # Prepare for import
        import_data = []
        for kw in keywords_data:
            artist_name = kw.get('artist_name', '').strip()
            artist_id = artist_map.get(artist_name) if artist_name else None
            
            import_data.append({
                'keyword': kw['keyword'],
                'artist_id': artist_id,
                'auto_flag': kw.get('auto_flag', False),
                'priority': kw.get('priority', 'Medium')
            })
        
        # Bulk import
        result = db.bulk_add_keywords(import_data)
        return jsonify(result)
        
    except Exception as e:
        return jsonify({'error': f'Import failed: {str(e)}'}), 500

@app.route('/api/keywords/<int:keyword_id>', methods=['PUT'])
def update_keyword(keyword_id):
    data = request.json
    success = db.update_keyword(
        keyword_id,
        active=data.get('active'),
        auto_flag=data.get('auto_flag'),
        priority=data.get('priority')
    )
    
    if success:
        return jsonify({'success': True, 'message': 'Keyword updated'})
    else:
        return jsonify({'error': 'Keyword not found'}), 404

@app.route('/api/keywords/<int:keyword_id>', methods=['DELETE'])
def delete_keyword(keyword_id):
    success = db.delete_keyword(keyword_id)
    
    if success:
        return jsonify({'success': True, 'message': 'Keyword deleted'})
    else:
        return jsonify({'error': 'Keyword not found'}), 404

# ============================================================================
# AUTO-FLAG RULES ENDPOINTS
# ============================================================================

@app.route('/api/auto-flag-rules', methods=['GET'])
def get_auto_flag_rules():
    rules = db.get_active_auto_flag_rules()
    # Also get inactive rules
    all_rules = []
    conn = db.get_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM auto_flag_rules ORDER BY created_at DESC")
    rows = cursor.fetchall()
    conn.close()
    
    from models import AutoFlagRule
    for row in rows:
        rule = AutoFlagRule(
            id=row['id'],
            name=row['name'],
            description=row['description'],
            conditions=row['conditions'],
            action=row['action'],
            active=bool(row['active']),
            created_at=row['created_at']
        )
        all_rules.append(rule.to_dict())
    
    return jsonify(all_rules)

@app.route('/api/auto-flag-rules', methods=['POST'])
def add_auto_flag_rule():
    data = request.json
    name = data.get('name', '').strip()
    description = data.get('description', '').strip()
    conditions = data.get('conditions', {})
    action = data.get('action', 'flag')
    
    if not name:
        return jsonify({'error': 'Rule name is required'}), 400
    
    if not conditions:
        return jsonify({'error': 'At least one condition is required'}), 400
    
    rule_id = db.add_auto_flag_rule(name, conditions, action, description)
    
    if rule_id:
        return jsonify({'success': True, 'id': rule_id, 'message': 'Rule created'})
    else:
        return jsonify({'error': 'Rule with this name already exists'}), 400

@app.route('/api/auto-flag-rules/<int:rule_id>', methods=['PUT'])
def update_auto_flag_rule(rule_id):
    data = request.json
    active = data.get('active')
    
    if active is None:
        return jsonify({'error': 'Active status is required'}), 400
    
    conn = db.get_connection()
    cursor = conn.cursor()
    cursor.execute("UPDATE auto_flag_rules SET active = ? WHERE id = ?", (int(active), rule_id))
    conn.commit()
    success = cursor.rowcount > 0
    conn.close()
    
    if success:
        return jsonify({'success': True, 'message': 'Rule updated'})
    else:
        return jsonify({'error': 'Rule not found'}), 404

@app.route('/api/auto-flag-rules/<int:rule_id>', methods=['DELETE'])
def delete_auto_flag_rule(rule_id):
    conn = db.get_connection()
    cursor = conn.cursor()
    cursor.execute("DELETE FROM auto_flag_rules WHERE id = ?", (rule_id,))
    conn.commit()
    success = cursor.rowcount > 0
    conn.close()
    
    if success:
        return jsonify({'success': True, 'message': 'Rule deleted'})
    else:
        return jsonify({'error': 'Rule not found'}), 404

# ============================================================================
# SEARCH ENDPOINT (Enhanced with auto-flagging)
# ============================================================================

@app.route('/api/search', methods=['POST'])
def run_search():
    if not API_CONFIGURED or not youtube_service:
        return jsonify({'error': 'YouTube API is not configured'}), 500
    
    data = request.json
    keywords = data.get('keywords', [])
    
    if not keywords:
        keywords = db.get_active_keywords()
    
    if not keywords:
        return jsonify({'error': 'No keywords to search'}), 400
    
    results = {
        'total_found': 0,
        'total_new': 0,
        'keywords': []
    }
    
    for keyword in keywords:
        try:
            videos = youtube_service.search_videos(keyword)
            new_count = 0
            
            for video in videos:
                # Apply auto-flag rules
                should_flag, priority = db.apply_auto_flag_rules(video)
                
                video.auto_flagged = should_flag
                video.priority = priority
                if should_flag:
                    video.status = 'Flagged for Takedown'
                
                video_id = db.add_video(video)
                if video_id:
                    new_count += 1
                    
                    # Send email alert for critical
                    if priority == 'Critical' and email_service.enabled:
                        # Get artist email if available
                        if video.artist_id:
                            artist = db.get_artist(video.artist_id)
                            if artist and artist.email:
                                email_service.send_critical_alert(
                                    artist.email,
                                    video.to_dict(),
                                    "Auto-flagged as Critical priority"
                                )
            
            results['total_found'] += len(videos)
            results['total_new'] += new_count
            results['keywords'].append({
                'keyword': keyword,
                'found': len(videos),
                'new': new_count
            })
            
            db.add_search_log(keyword, len(videos), True)
            
        except Exception as e:
            error_msg = str(e)
            results['keywords'].append({
                'keyword': keyword,
                'error': error_msg
            })
            db.add_search_log(keyword, 0, False, error_msg)
    
    return jsonify(results)

# ============================================================================
# VIDEOS ENDPOINTS (Enhanced)
# ============================================================================

@app.route('/api/videos', methods=['GET'])
def get_videos():
    filters = {}
    
    if request.args.get('keyword'):
        filters['keyword'] = request.args.get('keyword')
    if request.args.get('status'):
        filters['status'] = request.args.get('status')
    if request.args.get('priority'):
        filters['priority'] = request.args.get('priority')
    if request.args.get('artist_id'):
        filters['artist_id'] = int(request.args.get('artist_id'))
    if request.args.get('auto_flagged'):
        filters['auto_flagged'] = request.args.get('auto_flagged').lower() == 'true'
    if request.args.get('date_from'):
        filters['date_from'] = request.args.get('date_from')
    if request.args.get('date_to'):
        filters['date_to'] = request.args.get('date_to')
    
    videos = db.get_all_videos(filters)
    return jsonify([v.to_dict() for v in videos])

@app.route('/api/videos/<int:video_id>', methods=['PUT'])
def update_video(video_id):
    data = request.json
    status = data.get('status')
    priority = data.get('priority')
    
    conn = db.get_connection()
    cursor = conn.cursor()
    
    if status:
        if status not in ['Pending', 'Reviewed', 'Flagged for Takedown']:
            return jsonify({'error': 'Invalid status'}), 400
        cursor.execute("UPDATE videos SET status = ? WHERE id = ?", (status, video_id))
    
    if priority:
        if priority not in ['Low', 'Medium', 'High', 'Critical']:
            return jsonify({'error': 'Invalid priority'}), 400
        cursor.execute("UPDATE videos SET priority = ? WHERE id = ?", (priority, video_id))
    
    conn.commit()
    success = cursor.rowcount > 0
    conn.close()
    
    if success:
        return jsonify({'success': True, 'message': 'Video updated'})
    else:
        return jsonify({'error': 'Video not found'}), 404

@app.route('/api/videos/<int:video_id>', methods=['DELETE'])
def delete_video(video_id):
    success = db.delete_video(video_id)
    
    if success:
        return jsonify({'success': True, 'message': 'Video deleted'})
    else:
        return jsonify({'error': 'Video not found'}), 404

@app.route('/api/videos/batch-update', methods=['POST'])
def batch_update_videos():
    data = request.json
    video_ids = data.get('video_ids', [])
    status = data.get('status')
    priority = data.get('priority')
    
    if not video_ids:
        return jsonify({'error': 'No videos selected'}), 400
    
    count = db.batch_update_videos(video_ids, status, priority)
    
    return jsonify({
        'success': True,
        'updated': count,
        'message': f'{count} videos updated'
    })

@app.route('/api/videos/batch-delete', methods=['POST'])
def batch_delete_videos():
    data = request.json
    video_ids = data.get('video_ids', [])
    
    if not video_ids:
        return jsonify({'error': 'No videos selected'}), 400
    
    conn = db.get_connection()
    cursor = conn.cursor()
    
    placeholders = ','.join('?' * len(video_ids))
    cursor.execute(f"DELETE FROM videos WHERE id IN ({placeholders})", video_ids)
    conn.commit()
    count = cursor.rowcount
    conn.close()
    
    return jsonify({
        'success': True,
        'deleted': count,
        'message': f'{count} videos deleted'
    })

# ============================================================================
# STATISTICS ENDPOINT (Enhanced)
# ============================================================================

@app.route('/api/stats', methods=['GET'])
def get_stats():
    artist_id = request.args.get('artist_id', type=int)
    stats = db.get_stats(artist_id)
    return jsonify(stats)

# ============================================================================
# SEARCH LOGS ENDPOINT
# ============================================================================

@app.route('/api/logs', methods=['GET'])
def get_logs():
    limit = request.args.get('limit', 50, type=int)
    artist_id = request.args.get('artist_id', type=int)
    logs = db.get_search_logs(limit, artist_id)
    return jsonify([log.to_dict() for log in logs])

# ============================================================================
# EXPORT ENDPOINTS
# ============================================================================

@app.route('/api/export/csv', methods=['GET'])
def export_csv():
    filters = {}
    if request.args.get('keyword'):
        filters['keyword'] = request.args.get('keyword')
    if request.args.get('status'):
        filters['status'] = request.args.get('status')
    if request.args.get('priority'):
        filters['priority'] = request.args.get('priority')
    if request.args.get('artist_id'):
        filters['artist_id'] = int(request.args.get('artist_id'))
    
    videos = db.get_all_videos(filters)
    
    output = io.StringIO()
    writer = csv.writer(output)
    
    writer.writerow(['ID', 'Video ID', 'Title', 'Channel Name', 'Channel ID', 
                     'Publish Date', 'Video URL', 'Matched Keyword', 'Status', 
                     'Priority', 'Auto-Flagged', 'Created At'])
    
    for video in videos:
        writer.writerow([
            video.id, video.video_id, video.title, video.channel_name,
            video.channel_id, video.publish_date, video.video_url,
            video.matched_keyword, video.status, video.priority,
            'Yes' if video.auto_flagged else 'No', video.created_at
        ])
    
    output.seek(0)
    
    return send_file(
        io.BytesIO(output.getvalue().encode('utf-8')),
        mimetype='text/csv',
        as_attachment=True,
        download_name=f'ugc_videos_{datetime.now().strftime("%Y%m%d_%H%M%S")}.csv'
    )

@app.route('/api/export/excel', methods=['GET'])
def export_excel():
    filters = {}
    if request.args.get('keyword'):
        filters['keyword'] = request.args.get('keyword')
    if request.args.get('status'):
        filters['status'] = request.args.get('status')
    if request.args.get('priority'):
        filters['priority'] = request.args.get('priority')
    if request.args.get('artist_id'):
        filters['artist_id'] = int(request.args.get('artist_id'))
    
    videos = db.get_all_videos(filters)
    
    wb = Workbook()
    ws = wb.active
    ws.title = "UGC Videos"
    
    headers = ['ID', 'Video ID', 'Title', 'Channel Name', 'Channel ID', 
               'Publish Date', 'Video URL', 'Matched Keyword', 'Status',
               'Priority', 'Auto-Flagged', 'Created At']
    ws.append(headers)
    
    for video in videos:
        ws.append([
            video.id, video.video_id, video.title, video.channel_name,
            video.channel_id, video.publish_date, video.video_url,
            video.matched_keyword, video.status, video.priority,
            'Yes' if video.auto_flagged else 'No', video.created_at
        ])
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=f'ugc_videos_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
    )

# ============================================================================
# SCHEDULER ENDPOINT
# ============================================================================

@app.route('/api/schedule', methods=['POST'])
def setup_schedule():
    data = request.json
    enabled = data.get('enabled', False)
    interval_hours = data.get('interval_hours', 24)
    
    if enabled:
        for job in scheduler.get_jobs():
            job.remove()
        
        def scheduled_search():
            with app.app_context():
                keywords = db.get_active_keywords()
                for keyword in keywords:
                    try:
                        videos = youtube_service.search_videos(keyword)
                        for video in videos:
                            # Apply auto-flag rules
                            should_flag, priority = db.apply_auto_flag_rules(video)
                            video.auto_flagged = should_flag
                            video.priority = priority
                            if should_flag:
                                video.status = 'Flagged for Takedown'
                            db.add_video(video)
                        db.add_search_log(keyword, len(videos), True)
                    except Exception as e:
                        db.add_search_log(keyword, 0, False, str(e))
        
        scheduler.add_job(
            scheduled_search,
            'interval',
            hours=interval_hours,
            id='auto_search'
        )
        
        return jsonify({'success': True, 'message': f'Automatic search scheduled every {interval_hours} hours'})
    else:
        for job in scheduler.get_jobs():
            job.remove()
        return jsonify({'success': True, 'message': 'Automatic search disabled'})

# ============================================================================
# NOTIFICATIONS ENDPOINTS
# ============================================================================

@app.route('/api/notifications/test', methods=['POST'])
def test_notification():
    data = request.json
    email = data.get('email')
    
    if not email:
        return jsonify({'error': 'Email is required'}), 400
    
    if not email_service.enabled:
        return jsonify({'error': 'Email service not configured'}), 500
    
    success = email_service.send_email(
        email,
        'Test Email from UGC Monitor',
        'This is a test email. Your email notifications are working!'
    )
    
    if success:
        return jsonify({'success': True, 'message': 'Test email sent'})
    else:
        return jsonify({'error': 'Failed to send test email'}), 500

if __name__ == '__main__':
    port = int(os.getenv('PORT', 5000))
    app.run(host='0.0.0.0', port=port, debug=True)
