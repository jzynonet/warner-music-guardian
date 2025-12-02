from flask import Flask, request, jsonify, send_file
from flask_cors import CORS
from dotenv import load_dotenv
from datetime import datetime
from apscheduler.schedulers.background import BackgroundScheduler
import os
import csv
import io
from openpyxl import Workbook

from database_enhanced import Database
from youtube_service import YouTubeService
from models import Video
from bulk_import import BulkImporter
from email_service import EmailService
from music_detector import MusicDetector, analyze_video_for_piracy, get_smart_rules
from spotify_service import SpotifyService
from musicbrainz_service import MusicBrainzService
from keyword_learning import KeywordLearning
from auto_update_service import AutoUpdateService

# Load environment variables
load_dotenv()

app = Flask(__name__)
app.config['SECRET_KEY'] = os.getenv('FLASK_SECRET_KEY', 'dev-secret-key')

# Configure CORS to allow frontend access
CORS(app, resources={
    r"/api/*": {
        "origins": ["http://localhost:5173", "http://localhost:3000", "http://127.0.0.1:5173"],
        "methods": ["GET", "POST", "PUT", "DELETE", "OPTIONS"],
        "allow_headers": ["Content-Type", "Authorization"],
        "supports_credentials": True
    }
})

# Initialize database
db = Database(os.getenv('DATABASE_PATH', 'videos.db'))

# Initialize YouTube service
try:
    youtube_service = YouTubeService(os.getenv('YOUTUBE_API_KEY', ''))
    API_CONFIGURED = True
except ValueError as e:
    youtube_service = None
    API_CONFIGURED = False
    print(f"WARNING: {str(e)}")

# Initialize email service
email_service = EmailService()

# Initialize music detector
music_detector = MusicDetector()

# Initialize Spotify service (optional)
try:
    spotify_service = SpotifyService()
    SPOTIFY_CONFIGURED = spotify_service.test_credentials()
except:
    spotify_service = None
    SPOTIFY_CONFIGURED = False

# Initialize MusicBrainz service (always available, no auth needed)
musicbrainz_service = MusicBrainzService()

# Initialize keyword learning
keyword_learner = KeywordLearning(db)

# Initialize auto-update service
auto_update_service = AutoUpdateService(db, spotify_service, musicbrainz_service)

# Admin password
ADMIN_PASSWORD = os.getenv('ADMIN_PASSWORD', 'admin123')

# Debug: Print loaded password on startup
print("=" * 50)
print(f"ADMIN PASSWORD LOADED: '{ADMIN_PASSWORD}'")
print(f"EMAIL SERVICE: {'Enabled' if email_service.enabled else 'Disabled'}")
print("=" * 50)

# Scheduler for automatic searches
scheduler = BackgroundScheduler()
scheduler.start()

# Auth middleware
def check_auth(password):
    return password == ADMIN_PASSWORD

@app.route('/api/health', methods=['GET'])
def health_check():
    return jsonify({
        'status': 'ok',
        'api_configured': API_CONFIGURED,
        'database_ok': os.path.exists(db.db_path),
        'email_configured': email_service.enabled,
        'spotify_configured': SPOTIFY_CONFIGURED,
        'musicbrainz_available': True
    })

@app.route('/api/auth/login', methods=['POST'])
def login():
    data = request.json
    password = data.get('password', '')
    
    # Debug logging
    print(f"Login attempt - Received password: '{password}'")
    print(f"Expected password: '{ADMIN_PASSWORD}'")
    print(f"Passwords match: {password == ADMIN_PASSWORD}")
    
    if check_auth(password):
        return jsonify({'success': True, 'message': 'Authentication successful'})
    else:
        return jsonify({'success': False, 'message': 'Invalid password'}), 401

# ============================================================================
# ARTISTS ENDPOINTS
# ============================================================================

@app.route('/api/artists', methods=['GET'])
def get_artists():
    artists = db.get_all_artists()
    return jsonify([a.to_dict() for a in artists])

@app.route('/api/artists', methods=['POST'])
def add_artist():
    data = request.json
    name = data.get('name', '').strip()
    email = data.get('email', '').strip()
    contact_person = data.get('contact_person', '').strip()
    notes = data.get('notes', '').strip()
    
    if not name:
        return jsonify({'error': 'Artist name is required'}), 400
    
    artist_id = db.add_artist(name, email, contact_person, notes)
    
    if artist_id:
        return jsonify({'success': True, 'id': artist_id, 'message': 'Artist added'})
    else:
        return jsonify({'error': 'Artist already exists'}), 400

@app.route('/api/artists/<int:artist_id>', methods=['GET'])
def get_artist(artist_id):
    artist = db.get_artist(artist_id)
    if artist:
        return jsonify(artist.to_dict())
    else:
        return jsonify({'error': 'Artist not found'}), 404

@app.route('/api/artists/<int:artist_id>', methods=['PUT'])
def update_artist(artist_id):
    data = request.json
    success = db.update_artist(
        artist_id,
        name=data.get('name'),
        email=data.get('email'),
        contact_person=data.get('contact_person'),
        notes=data.get('notes'),
        active=data.get('active')
    )
    
    if success:
        return jsonify({'success': True, 'message': 'Artist updated'})
    else:
        return jsonify({'error': 'Artist not found'}), 404

@app.route('/api/artists/<int:artist_id>', methods=['DELETE'])
def delete_artist(artist_id):
    success = db.delete_artist(artist_id)
    
    if success:
        return jsonify({'success': True, 'message': 'Artist deleted'})
    else:
        return jsonify({'error': 'Artist not found'}), 404

# ============================================================================
# KEYWORDS ENDPOINTS (Enhanced)
# ============================================================================

@app.route('/api/keywords', methods=['GET'])
def get_keywords():
    artist_id = request.args.get('artist_id', type=int)
    keywords = db.get_all_keywords(artist_id)
    return jsonify([k.to_dict() for k in keywords])

@app.route('/api/keywords', methods=['POST'])
def add_keyword():
    data = request.json
    keyword = data.get('keyword', '').strip()
    artist_id = data.get('artist_id')
    auto_flag = data.get('auto_flag', False)
    priority = data.get('priority', 'Medium')
    
    if not keyword:
        return jsonify({'error': 'Keyword is required'}), 400
    
    keyword_id = db.add_keyword(keyword, artist_id, auto_flag, priority)
    
    if keyword_id:
        return jsonify({'success': True, 'id': keyword_id, 'message': 'Keyword added'})
    else:
        return jsonify({'error': 'Keyword already exists'}), 400

@app.route('/api/keywords/bulk-import', methods=['POST'])
def bulk_import_keywords():
    if 'file' not in request.files:
        return jsonify({'error': 'No file provided'}), 400
    
    file = request.files['file']
    if file.filename == '':
        return jsonify({'error': 'No file selected'}), 400
    
    try:
        # Read file content
        if file.filename.endswith('.csv'):
            content = file.read().decode('utf-8')
            keywords_data = BulkImporter.parse_csv(content)
        elif file.filename.endswith('.xlsx'):
            # Save temporarily
            temp_path = f"/tmp/{file.filename}"
            file.save(temp_path)
            keywords_data = BulkImporter.parse_excel(temp_path)
            os.remove(temp_path)
        else:
            return jsonify({'error': 'Invalid file format. Use CSV or XLSX'}), 400
        
        # Get artist mapping
        artists = db.get_all_artists()
        artist_map = {a.name: a.id for a in artists}
        
        # Prepare for import
        import_data = []
        for kw in keywords_data:
            artist_name = kw.get('artist_name', '').strip()
            artist_id = artist_map.get(artist_name) if artist_name else None
            
            import_data.append({
                'keyword': kw['keyword'],
                'artist_id': artist_id,
                'auto_flag': kw.get('auto_flag', False),
                'priority': kw.get('priority', 'Medium')
            })
        
        # Bulk import
        result = db.bulk_add_keywords(import_data)
        return jsonify(result)
        
    except Exception as e:
        return jsonify({'error': f'Import failed: {str(e)}'}), 500

@app.route('/api/keywords/<int:keyword_id>', methods=['PUT'])
def update_keyword(keyword_id):
    data = request.json
    success = db.update_keyword(
        keyword_id,
        active=data.get('active'),
        auto_flag=data.get('auto_flag'),
        priority=data.get('priority')
    )
    
    if success:
        return jsonify({'success': True, 'message': 'Keyword updated'})
    else:
        return jsonify({'error': 'Keyword not found'}), 404

@app.route('/api/keywords/<int:keyword_id>', methods=['DELETE'])
def delete_keyword(keyword_id):
    success = db.delete_keyword(keyword_id)
    
    if success:
        return jsonify({'success': True, 'message': 'Keyword deleted'})
    else:
        return jsonify({'error': 'Keyword not found'}), 404

@app.route('/api/keywords/clear', methods=['DELETE'])
def clear_all_keywords():
    """Delete all keywords"""
    try:
        artist_id = request.args.get('artist_id', type=int)
        conn = db.get_connection()
        cursor = conn.cursor()
        
        if artist_id:
            cursor.execute("DELETE FROM keywords WHERE artist_id = ?", (artist_id,))
        else:
            cursor.execute("DELETE FROM keywords")
        
        deleted_count = cursor.rowcount
        conn.commit()
        conn.close()
        
        return jsonify({
            'success': True, 
            'message': f'{deleted_count} keyword(s) cleared',
            'count': deleted_count
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# ============================================================================
# SONGS ENDPOINTS (Song + Artist Combinations)
# ============================================================================

@app.route('/api/songs', methods=['GET'])
def get_songs():
    artist_id = request.args.get('artist_id', type=int)
    songs = db.get_all_songs(artist_id)
    return jsonify([s.to_dict() for s in songs])

@app.route('/api/songs', methods=['POST'])
def add_song():
    data = request.json
    song_name = data.get('song_name', '').strip()
    artist_name = data.get('artist_name', '').strip()
    artist_id = data.get('artist_id')
    auto_flag = data.get('auto_flag', False)
    priority = data.get('priority', 'Medium')
    duration_ms = data.get('duration_ms')
    
    if not song_name or not artist_name:
        return jsonify({'error': 'Song name and artist name are required'}), 400
    
    song_id = db.add_song(song_name, artist_name, artist_id, auto_flag, priority, duration_ms)
    
    if song_id:
        return jsonify({'success': True, 'id': song_id, 'message': 'Song added'})
    else:
        return jsonify({'error': 'Song already exists'}), 400

@app.route('/api/songs/<int:song_id>', methods=['PUT'])
def update_song(song_id):
    data = request.json
    success = db.update_song(
        song_id,
        active=data.get('active'),
        auto_flag=data.get('auto_flag'),
        priority=data.get('priority')
    )
    
    if success:
        return jsonify({'success': True, 'message': 'Song updated'})
    else:
        return jsonify({'error': 'Song not found'}), 404

@app.route('/api/songs/<int:song_id>', methods=['DELETE'])
def delete_song(song_id):
    success = db.delete_song(song_id)
    
    if success:
        return jsonify({'success': True, 'message': 'Song deleted'})
    else:
        return jsonify({'error': 'Song not found'}), 404

@app.route('/api/songs/clear', methods=['DELETE'])
def clear_all_songs():
    """Delete all songs"""
    try:
        artist_id = request.args.get('artist_id', type=int)
        conn = db.get_connection()
        cursor = conn.cursor()
        
        if artist_id:
            cursor.execute("DELETE FROM songs WHERE artist_id = ?", (artist_id,))
        else:
            cursor.execute("DELETE FROM songs")
        
        deleted_count = cursor.rowcount
        conn.commit()
        conn.close()
        
        return jsonify({
            'success': True, 
            'message': f'{deleted_count} song(s) cleared',
            'count': deleted_count
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/songs/bulk-import', methods=['POST'])
def bulk_import_songs():
    if 'file' not in request.files:
        return jsonify({'error': 'No file provided'}), 400
    
    file = request.files['file']
    if file.filename == '':
        return jsonify({'error': 'No file selected'}), 400
    
    try:
        import csv
        import io
        
        content = file.read().decode('utf-8')
        csv_reader = csv.DictReader(io.StringIO(content))
        
        songs_to_import = []
        for row in csv_reader:
            song_name = row.get('song_name', '').strip()
            artist_name = row.get('artist_name', '').strip()
            
            if song_name and artist_name:
                songs_to_import.append({
                    'song_name': song_name,
                    'artist_name': artist_name,
                    'active': row.get('active', 'true').lower() == 'true',
                    'artist_id': int(row['artist_id']) if row.get('artist_id') else None,
                    'auto_flag': row.get('auto_flag', 'false').lower() == 'true',
                    'priority': row.get('priority', 'Medium')
                })
        
        result = db.bulk_add_songs(songs_to_import)
        return jsonify(result)
        
    except Exception as e:
        return jsonify({'error': f'Import failed: {str(e)}'}), 500

@app.route('/api/songs/preview-from-spotify', methods=['POST'])
def preview_from_spotify():
    """Preview artist songs from Spotify before importing"""
    if not spotify_service or not SPOTIFY_CONFIGURED:
        return jsonify({'error': 'Spotify API not configured. Please set SPOTIFY_CLIENT_ID and SPOTIFY_CLIENT_SECRET in .env'}), 500
    
    try:
        data = request.json
        spotify_url = data.get('spotify_url', '').strip()
        
        if not spotify_url:
            return jsonify({'error': 'Spotify URL is required'}), 400
        
        # Fetch artist and songs from Spotify
        result = spotify_service.get_artist_all_songs_by_url(spotify_url)
        
        if 'error' in result:
            return jsonify({'error': result['error']}), 400
        
        return jsonify({
            'success': True,
            'artist_info': result['artist_info'],
            'main_songs': result['main_songs'],
            'featured_songs': result['featured_songs'],
            'total_main_songs': result['total_main_songs'],
            'total_featured_songs': result['total_featured_songs'],
            'albums': len(result['albums'])
        })
        
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/songs/import-from-spotify', methods=['POST'])
def import_from_spotify():
    """Import selected songs from Spotify artist"""
    if not spotify_service or not SPOTIFY_CONFIGURED:
        return jsonify({'error': 'Spotify API not configured. Please set SPOTIFY_CLIENT_ID and SPOTIFY_CLIENT_SECRET in .env'}), 500
    
    try:
        data = request.json
        artist_info = data.get('artist_info')
        selected_songs = data.get('selected_songs', [])
        auto_flag = data.get('auto_flag', False)
        priority = data.get('priority', 'Medium')
        
        if not artist_info or not selected_songs:
            return jsonify({'error': 'Artist info and selected songs are required'}), 400
        
        # Check if artist already exists in database
        existing_artist = None
        all_artists = db.get_all_artists()
        for artist in all_artists:
            if artist.name.lower() == artist_info['name'].lower():
                existing_artist = artist
                break
        
        # Add artist to database if not exists
        if existing_artist:
            artist_id = existing_artist.id
        else:
            artist_id = db.add_artist(
                name=artist_info['name'],
                notes=f"Imported from Spotify. Followers: {artist_info.get('followers', 0):,}"
            )
        
        # Add selected songs to database
        songs_to_import = []
        for song_data in selected_songs:
            # Handle both string and dict formats
            if isinstance(song_data, dict):
                song_name = song_data['name']
                duration_ms = song_data.get('duration_ms')
            else:
                song_name = song_data
                duration_ms = None
                
            songs_to_import.append({
                'song_name': song_name,
                'artist_name': artist_info['name'],
                'active': True,
                'artist_id': artist_id,
                'auto_flag': auto_flag,
                'priority': priority,
                'duration_ms': duration_ms
            })
        
        import_result = db.bulk_add_songs(songs_to_import)
        
        return jsonify({
            'success': True,
            'artist': {
                'id': artist_id,
                'name': artist_info['name'],
                'spotify_url': artist_info.get('spotify_url'),
                'followers': artist_info.get('followers'),
                'genres': artist_info.get('genres', [])
            },
            'songs_added': import_result['added'],
            'songs_skipped': import_result['skipped'],
            'total_songs_selected': len(selected_songs),
            'message': f"Successfully imported {import_result['added']} songs from {artist_info['name']}"
        })
        
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# ============================================================================
# AUTO-FLAG RULES ENDPOINTS
# ============================================================================

@app.route('/api/auto-flag-rules', methods=['GET'])
def get_auto_flag_rules():
    rules = db.get_active_auto_flag_rules()
    # Also get inactive rules
    all_rules = []
    conn = db.get_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM auto_flag_rules ORDER BY created_at DESC")
    rows = cursor.fetchall()
    conn.close()
    
    from models import AutoFlagRule
    for row in rows:
        rule = AutoFlagRule(
            id=row['id'],
            name=row['name'],
            description=row['description'],
            conditions=row['conditions'],
            action=row['action'],
            active=bool(row['active']),
            created_at=row['created_at']
        )
        all_rules.append(rule.to_dict())
    
    return jsonify(all_rules)

@app.route('/api/auto-flag-rules', methods=['POST'])
def add_auto_flag_rule():
    data = request.json
    name = data.get('name', '').strip()
    description = data.get('description', '').strip()
    conditions = data.get('conditions', {})
    action = data.get('action', 'flag')
    
    if not name:
        return jsonify({'error': 'Rule name is required'}), 400
    
    if not conditions:
        return jsonify({'error': 'At least one condition is required'}), 400
    
    rule_id = db.add_auto_flag_rule(name, conditions, action, description)
    
    if rule_id:
        return jsonify({'success': True, 'id': rule_id, 'message': 'Rule created'})
    else:
        return jsonify({'error': 'Rule with this name already exists'}), 400

@app.route('/api/auto-flag-rules/<int:rule_id>', methods=['PUT'])
def update_auto_flag_rule(rule_id):
    data = request.json
    active = data.get('active')
    
    if active is None:
        return jsonify({'error': 'Active status is required'}), 400
    
    conn = db.get_connection()
    cursor = conn.cursor()
    cursor.execute("UPDATE auto_flag_rules SET active = ? WHERE id = ?", (int(active), rule_id))
    conn.commit()
    success = cursor.rowcount > 0
    conn.close()
    
    if success:
        return jsonify({'success': True, 'message': 'Rule updated'})
    else:
        return jsonify({'error': 'Rule not found'}), 404

@app.route('/api/auto-flag-rules/<int:rule_id>', methods=['DELETE'])
def delete_auto_flag_rule(rule_id):
    conn = db.get_connection()
    cursor = conn.cursor()
    cursor.execute("DELETE FROM auto_flag_rules WHERE id = ?", (rule_id,))
    conn.commit()
    success = cursor.rowcount > 0
    conn.close()
    
    if success:
        return jsonify({'success': True, 'message': 'Rule deleted'})
    else:
        return jsonify({'error': 'Rule not found'}), 404

# ============================================================================
# SMART MUSIC DETECTION ENDPOINTS
# ============================================================================

@app.route('/api/smart-rules', methods=['GET'])
def get_smart_detection_rules():
    """Get recommended pre-built smart detection rules"""
    try:
        smart_rules = get_smart_rules()
        return jsonify({
            'success': True,
            'rules': smart_rules,
            'count': len(smart_rules)
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/smart-rules/install', methods=['POST'])
def install_smart_rules():
    """Install all pre-built smart detection rules"""
    try:
        smart_rules = get_smart_rules()
        installed_count = 0
        skipped_count = 0
        errors = []
        
        for rule in smart_rules:
            try:
                # Check if rule already exists
                existing = db.get_auto_flag_rule_by_name(rule['name'])
                if existing:
                    skipped_count += 1
                    continue
                
                # Add the rule
                rule_id = db.add_auto_flag_rule(
                    name=rule['name'],
                    description=rule['description'],
                    conditions=rule['conditions'],
                    action=rule['action'],
                    active=rule.get('active', True)
                )
                installed_count += 1
                
            except Exception as e:
                errors.append(f"{rule['name']}: {str(e)}")
        
        return jsonify({
            'success': True,
            'installed': installed_count,
            'skipped': skipped_count,
            'errors': errors,
            'message': f'Installed {installed_count} smart rules, skipped {skipped_count} existing'
        })
        
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/analyze-video', methods=['POST'])
def analyze_video_ai():
    """
    Analyze a specific video using AI-powered music detection
    Request body: { video_id: str } or { title: str, channel_name: str, description: str }
    """
    try:
        data = request.json
        
        # Option 1: Analyze existing video by ID
        if 'video_id' in data:
            video_id = data['video_id']
            video = db.get_video_by_id(video_id)
            if not video:
                return jsonify({'error': 'Video not found'}), 404
            
            analysis = music_detector.analyze_video(
                title=video.title,
                channel_name=video.channel_name,
                description=getattr(video, 'description', ''),
                view_count=getattr(video, 'view_count', 0),
                duration=0  # Could add duration to Video model
            )
            
            return jsonify({
                'success': True,
                'video_id': video_id,
                'video_title': video.title,
                'analysis': analysis
            })
        
        # Option 2: Analyze provided video data
        elif 'title' in data and 'channel_name' in data:
            analysis = music_detector.analyze_video(
                title=data['title'],
                channel_name=data['channel_name'],
                description=data.get('description', ''),
                view_count=data.get('view_count', 0),
                duration=data.get('duration', 0)
            )
            
            return jsonify({
                'success': True,
                'analysis': analysis
            })
        
        else:
            return jsonify({'error': 'Either video_id or (title + channel_name) required'}), 400
            
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/smart-scan', methods=['POST'])
def smart_scan_all_videos():
    """
    Re-scan all videos with AI detection and update flags
    Optional: { artist_id: int } to scan only specific artist's videos
    """
    try:
        data = request.json or {}
        artist_id = data.get('artist_id')
        
        # Get videos to scan
        conn = db.get_connection()
        cursor = conn.cursor()
        
        if artist_id:
            cursor.execute("""
                SELECT id, title, channel_name, status
                FROM videos 
                WHERE artist_id = ?
            """, (artist_id,))
        else:
            cursor.execute("""
                SELECT id, title, channel_name, status
                FROM videos
            """)
        
        videos = cursor.fetchall()
        
        updated_count = 0
        flagged_count = 0
        
        for video in videos:
            video_id, title, channel_name, current_status = video
            
            # Analyze with AI
            analysis = music_detector.analyze_video(
                title=title,
                channel_name=channel_name
            )
            
            # Update if should flag and not already flagged
            if analysis['should_flag'] and current_status != 'Flagged for Takedown':
                cursor.execute("""
                    UPDATE videos 
                    SET status = 'Flagged for Takedown',
                        auto_flagged = 1
                    WHERE id = ?
                """, (video_id,))
                updated_count += 1
                flagged_count += 1
        
        conn.commit()
        conn.close()
        
        return jsonify({
            'success': True,
            'scanned': len(videos),
            'updated': updated_count,
            'newly_flagged': flagged_count,
            'message': f'Smart scan complete: {updated_count} videos updated'
        })
        
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# ============================================================================
# SEARCH ENDPOINT (Enhanced with auto-flagging)
# ============================================================================

@app.route('/api/search', methods=['POST'])
def run_search():
    if not API_CONFIGURED or not youtube_service:
        return jsonify({'error': 'YouTube API is not configured'}), 500
    
    data = request.json
    keywords = data.get('keywords', [])
    exclude_keywords = data.get('exclude_keywords', [])
    
    if not keywords:
        keywords = db.get_active_keywords()
    
    if not keywords:
        return jsonify({'error': 'No keywords to search'}), 400
    
    # Filter out excluded keywords
    if exclude_keywords:
        if isinstance(keywords[0], str):
            keywords = [k for k in keywords if k not in exclude_keywords]
        else:
            # If keywords are objects with a 'keyword' attribute
            keywords = [k for k in keywords if k.keyword not in exclude_keywords]
    
    results = {
        'total_found': 0,
        'total_new': 0,
        'keywords': []
    }
    
    for keyword in keywords:
        try:
            videos = youtube_service.search_videos(keyword)
            new_count = 0
            
            for video in videos:
                # Apply AI smart detection to EVERY video automatically
                ai_analysis = music_detector.analyze_video(
                    title=video.title,
                    channel_name=video.channel_name,
                    description=getattr(video, 'description', ''),
                    view_count=getattr(video, 'view_count', 0),
                    duration=getattr(video, 'duration', 0)
                )
                
                # Apply auto-flag rules (legacy system)
                should_flag, priority = db.apply_auto_flag_rules(video)
                
                # Use AI decision if it's more strict
                if ai_analysis['should_flag'] and not should_flag:
                    should_flag = True
                    
                # Set priority based on AI risk level
                ai_priority_map = {
                    'critical': 'Critical',
                    'high': 'High',
                    'medium': 'Medium',
                    'low': 'Low'
                }
                ai_priority = ai_priority_map.get(ai_analysis['risk_level'], 'Medium')
                
                # Use highest priority between AI and rules
                priority_rank = {'Critical': 4, 'High': 3, 'Medium': 2, 'Low': 1}
                if priority_rank.get(ai_priority, 0) > priority_rank.get(priority, 0):
                    priority = ai_priority
                
                video.auto_flagged = should_flag or ai_analysis['should_flag']
                video.priority = priority
                if video.auto_flagged:
                    video.status = 'Flagged for Takedown'
                
                # Store AI analysis results
                video.ai_risk_score = ai_analysis['risk_score']
                video.ai_risk_level = ai_analysis['risk_level']
                video.ai_reason = ai_analysis['reason']
                
                video_id = db.add_video(video)
                if video_id:
                    new_count += 1
                    
                    # Send email alert for critical
                    if priority == 'Critical' and email_service.enabled:
                        # Get artist email if available
                        if video.artist_id:
                            artist = db.get_artist(video.artist_id)
                            if artist and artist.email:
                                email_service.send_critical_alert(
                                    artist.email,
                                    video.to_dict(),
                                    "Auto-flagged as Critical priority"
                                )
            
            results['total_found'] += len(videos)
            results['total_new'] += new_count
            results['keywords'].append({
                'keyword': keyword,
                'found': len(videos),
                'new': new_count
            })
            
            db.add_search_log(keyword, len(videos), True)
            
        except Exception as e:
            error_msg = str(e)
            results['keywords'].append({
                'keyword': keyword,
                'error': error_msg
            })
            db.add_search_log(keyword, 0, False, error_msg)
    
    return jsonify(results)

@app.route('/api/search/songs', methods=['POST'])
def run_songs_search():
    """Search using song + artist combinations for more accurate results"""
    if not API_CONFIGURED or not youtube_service:
        return jsonify({'error': 'YouTube API is not configured'}), 500
    
    data = request.json
    songs = data.get('songs', [])
    
    if not songs:
        songs = db.get_active_songs()
    
    if not songs:
        return jsonify({'error': 'No songs to search'}), 400
    
    results = {
        'total_found': 0,
        'total_new': 0,
        'songs': []
    }
    
    for song in songs:
        try:
            # Extract song_name, artist_name, and duration_ms
            if isinstance(song, dict):
                song_name = song.get('song_name')
                artist_name = song.get('artist_name')
                duration_ms = song.get('duration_ms')
            else:
                song_name = song.song_name
                artist_name = song.artist_name
                duration_ms = song.duration_ms
            
            videos = youtube_service.search_song(song_name, artist_name, duration_ms=duration_ms)
            new_count = 0
            
            for video in videos:
                # Apply AI smart detection
                ai_analysis = music_detector.analyze_video(
                    title=video.title,
                    channel_name=video.channel_name,
                    description=getattr(video, 'description', ''),
                    view_count=getattr(video, 'view_count', 0),
                    duration=getattr(video, 'duration', 0)
                )
                
                # Apply auto-flag rules
                should_flag, priority = db.apply_auto_flag_rules(video)
                
                # Use AI decision if it's more strict
                if ai_analysis['should_flag'] and not should_flag:
                    should_flag = True
                    
                # Set priority based on AI risk level
                ai_priority_map = {
                    'critical': 'Critical',
                    'high': 'High',
                    'medium': 'Medium',
                    'low': 'Low'
                }
                ai_priority = ai_priority_map.get(ai_analysis['risk_level'], 'Medium')
                
                # Use highest priority
                priority_rank = {'Critical': 4, 'High': 3, 'Medium': 2, 'Low': 1}
                if priority_rank.get(ai_priority, 0) > priority_rank.get(priority, 0):
                    priority = ai_priority
                
                video.auto_flagged = should_flag or ai_analysis['should_flag']
                video.priority = priority
                if video.auto_flagged:
                    video.status = 'Flagged for Takedown'
                
                # Store AI analysis results
                video.ai_risk_score = ai_analysis['risk_score']
                video.ai_risk_level = ai_analysis['risk_level']
                video.ai_reason = ai_analysis['reason']
                
                video_id = db.add_video(video)
                if video_id:
                    new_count += 1
                    
                    # Send email alert for critical
                    if priority == 'Critical' and email_service.enabled:
                        if video.artist_id:
                            artist = db.get_artist(video.artist_id)
                            if artist and artist.email:
                                email_service.send_critical_alert(
                                    artist.email,
                                    video.to_dict(),
                                    "Auto-flagged as Critical priority"
                                )
            
            results['total_found'] += len(videos)
            results['total_new'] += new_count
            results['songs'].append({
                'song_name': song_name,
                'artist_name': artist_name,
                'found': len(videos),
                'new': new_count
            })
            
            # Log search with combined identifier
            db.add_search_log(f"{song_name} - {artist_name}", len(videos), True)
            
        except Exception as e:
            error_msg = str(e)
            song_name = song.get('song_name') if isinstance(song, dict) else song.song_name
            artist_name = song.get('artist_name') if isinstance(song, dict) else song.artist_name
            results['songs'].append({
                'song_name': song_name,
                'artist_name': artist_name,
                'error': error_msg
            })
            db.add_search_log(f"{song_name} - {artist_name}", 0, False, error_msg)
    
    return jsonify(results)

# ============================================================================
# VIDEOS ENDPOINTS (Enhanced)
# ============================================================================

@app.route('/api/videos', methods=['GET'])
def get_videos():
    filters = {}
    
    if request.args.get('keyword'):
        filters['keyword'] = request.args.get('keyword')
    if request.args.get('status'):
        filters['status'] = request.args.get('status')
    if request.args.get('priority'):
        filters['priority'] = request.args.get('priority')
    if request.args.get('artist_id'):
        filters['artist_id'] = int(request.args.get('artist_id'))
    if request.args.get('auto_flagged'):
        filters['auto_flagged'] = request.args.get('auto_flagged').lower() == 'true'
    if request.args.get('date_from'):
        filters['date_from'] = request.args.get('date_from')
    if request.args.get('date_to'):
        filters['date_to'] = request.args.get('date_to')
    
    videos = db.get_all_videos(filters)
    return jsonify([v.to_dict() for v in videos])

@app.route('/api/videos/<int:video_id>', methods=['PUT'])
def update_video(video_id):
    data = request.json
    status = data.get('status')
    priority = data.get('priority')
    
    conn = db.get_connection()
    cursor = conn.cursor()
    
    if status:
        if status not in ['Pending', 'Reviewed', 'Flagged for Takedown']:
            return jsonify({'error': 'Invalid status'}), 400
        cursor.execute("UPDATE videos SET status = ? WHERE id = ?", (status, video_id))
    
    if priority:
        if priority not in ['Low', 'Medium', 'High', 'Critical']:
            return jsonify({'error': 'Invalid priority'}), 400
        cursor.execute("UPDATE videos SET priority = ? WHERE id = ?", (priority, video_id))
    
    conn.commit()
    success = cursor.rowcount > 0
    conn.close()
    
    if success:
        return jsonify({'success': True, 'message': 'Video updated'})
    else:
        return jsonify({'error': 'Video not found'}), 404

@app.route('/api/videos/<int:video_id>', methods=['DELETE'])
def delete_video(video_id):
    success = db.delete_video(video_id)
    
    if success:
        return jsonify({'success': True, 'message': 'Video deleted'})
    else:
        return jsonify({'error': 'Video not found'}), 404

@app.route('/api/videos/batch-update', methods=['POST'])
def batch_update_videos():
    data = request.json
    video_ids = data.get('video_ids', [])
    status = data.get('status')
    priority = data.get('priority')
    
    if not video_ids:
        return jsonify({'error': 'No videos selected'}), 400
    
    count = db.batch_update_videos(video_ids, status, priority)
    
    return jsonify({
        'success': True,
        'updated': count,
        'message': f'{count} videos updated'
    })

@app.route('/api/videos/batch-delete', methods=['POST'])
def batch_delete_videos():
    data = request.json
    video_ids = data.get('video_ids', [])
    
    if not video_ids:
        return jsonify({'error': 'No videos selected'}), 400
    
    conn = db.get_connection()
    cursor = conn.cursor()
    
    placeholders = ','.join('?' * len(video_ids))
    cursor.execute(f"DELETE FROM videos WHERE id IN ({placeholders})", video_ids)
    conn.commit()
    count = cursor.rowcount
    conn.close()
    
    return jsonify({
        'success': True,
        'deleted': count,
        'message': f'{count} videos deleted'
    })

@app.route('/api/videos/clear-all', methods=['POST'])
def clear_all_videos():
    """Clear all videos from database"""
    try:
        import sqlite3
        conn = db.get_connection()
        
        # Get count before deletion
        cursor = conn.execute('SELECT COUNT(*) FROM videos')
        count = cursor.fetchone()[0]
        
        # Delete all videos and search logs
        conn.execute('DELETE FROM videos')
        conn.execute('DELETE FROM search_logs')
        conn.commit()
        
        return jsonify({
            'success': True, 
            'deleted': count,
            'message': f'Successfully deleted {count} videos'
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# ============================================================================
# STATISTICS ENDPOINT (Enhanced)
# ============================================================================

@app.route('/api/stats', methods=['GET'])
def get_stats():
    artist_id = request.args.get('artist_id', type=int)
    stats = db.get_stats(artist_id)
    return jsonify(stats)

# ============================================================================
# SEARCH LOGS ENDPOINT
# ============================================================================

@app.route('/api/logs', methods=['GET'])
def get_logs():
    limit = request.args.get('limit', 50, type=int)
    artist_id = request.args.get('artist_id', type=int)
    logs = db.get_search_logs(limit, artist_id)
    return jsonify([log.to_dict() for log in logs])

# ============================================================================
# EXPORT ENDPOINTS
# ============================================================================

@app.route('/api/export/csv', methods=['GET'])
def export_csv():
    filters = {}
    if request.args.get('keyword'):
        filters['keyword'] = request.args.get('keyword')
    if request.args.get('status'):
        filters['status'] = request.args.get('status')
    if request.args.get('priority'):
        filters['priority'] = request.args.get('priority')
    if request.args.get('artist_id'):
        filters['artist_id'] = int(request.args.get('artist_id'))
    
    videos = db.get_all_videos(filters)
    
    output = io.StringIO()
    writer = csv.writer(output)
    
    writer.writerow(['ID', 'Video ID', 'Title', 'Channel Name', 'Channel ID', 
                     'Publish Date', 'Video URL', 'Matched Keyword', 'Status', 
                     'Priority', 'Auto-Flagged', 'Created At'])
    
    for video in videos:
        writer.writerow([
            video.id, video.video_id, video.title, video.channel_name,
            video.channel_id, video.publish_date, video.video_url,
            video.matched_keyword, video.status, video.priority,
            'Yes' if video.auto_flagged else 'No', video.created_at
        ])
    
    output.seek(0)
    
    return send_file(
        io.BytesIO(output.getvalue().encode('utf-8')),
        mimetype='text/csv',
        as_attachment=True,
        download_name=f'ugc_videos_{datetime.now().strftime("%Y%m%d_%H%M%S")}.csv'
    )

@app.route('/api/export/excel', methods=['GET'])
def export_excel():
    filters = {}
    if request.args.get('keyword'):
        filters['keyword'] = request.args.get('keyword')
    if request.args.get('status'):
        filters['status'] = request.args.get('status')
    if request.args.get('priority'):
        filters['priority'] = request.args.get('priority')
    if request.args.get('artist_id'):
        filters['artist_id'] = int(request.args.get('artist_id'))
    
    videos = db.get_all_videos(filters)
    
    wb = Workbook()
    ws = wb.active
    ws.title = "UGC Videos"
    
    headers = ['ID', 'Video ID', 'Title', 'Channel Name', 'Channel ID', 
               'Publish Date', 'Video URL', 'Matched Keyword', 'Status',
               'Priority', 'Auto-Flagged', 'Created At']
    ws.append(headers)
    
    for video in videos:
        ws.append([
            video.id, video.video_id, video.title, video.channel_name,
            video.channel_id, video.publish_date, video.video_url,
            video.matched_keyword, video.status, video.priority,
            'Yes' if video.auto_flagged else 'No', video.created_at
        ])
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=f'ugc_videos_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
    )

# ============================================================================
# SCHEDULER ENDPOINT
# ============================================================================

@app.route('/api/schedule', methods=['POST'])
def setup_schedule():
    data = request.json
    enabled = data.get('enabled', False)
    interval_hours = data.get('interval_hours', 24)
    
    if enabled:
        for job in scheduler.get_jobs():
            job.remove()
        
        def scheduled_search():
            with app.app_context():
                keywords = db.get_active_keywords()
                for keyword in keywords:
                    try:
                        videos = youtube_service.search_videos(keyword)
                        for video in videos:
                            # Apply auto-flag rules
                            should_flag, priority = db.apply_auto_flag_rules(video)
                            video.auto_flagged = should_flag
                            video.priority = priority
                            if should_flag:
                                video.status = 'Flagged for Takedown'
                            db.add_video(video)
                        db.add_search_log(keyword, len(videos), True)
                    except Exception as e:
                        db.add_search_log(keyword, 0, False, str(e))
        
        scheduler.add_job(
            scheduled_search,
            'interval',
            hours=interval_hours,
            id='auto_search'
        )
        
        return jsonify({'success': True, 'message': f'Automatic search scheduled every {interval_hours} hours'})
    else:
        for job in scheduler.get_jobs():
            job.remove()
        return jsonify({'success': True, 'message': 'Automatic search disabled'})

# ============================================================================
# NOTIFICATIONS ENDPOINTS
# ============================================================================

@app.route('/api/notifications/test', methods=['POST'])
def test_notification():
    data = request.json
    email = data.get('email')
    
    if not email:
        return jsonify({'error': 'Email is required'}), 400
    
    if not email_service.enabled:
        return jsonify({'error': 'Email service not configured'}), 500
    
    success = email_service.send_email(
        email,
        'Test Email from UGC Monitor',
        'This is a test email. Your email notifications are working!'
    )
    
    if success:
        return jsonify({'success': True, 'message': 'Test email sent'})
    else:
        return jsonify({'error': 'Failed to send test email'}), 500

# ============================================================================
# KEYWORD AUTOMATION ENDPOINTS
# ============================================================================

@app.route('/api/keywords/fetch-spotify/<int:artist_id>', methods=['POST'])
def fetch_spotify_songs(artist_id):
    """Fetch all songs for an artist from Spotify"""
    if not spotify_service or not SPOTIFY_CONFIGURED:
        return jsonify({'error': 'Spotify API not configured'}), 500
    
    try:
        # Get artist
        artist = db.get_artist(artist_id)
        if not artist:
            return jsonify({'error': 'Artist not found'}), 404
        
        # Fetch songs from Spotify
        result = spotify_service.get_artist_all_songs(artist.name)
        
        if 'error' in result:
            return jsonify({'error': result['error']}), 500
        
        # Add songs as keywords
        added = 0
        skipped = 0
        errors = []
        
        for song in result['songs']:
            try:
                keyword_id = db.add_keyword(song, artist_id=artist_id, auto_flag=False, priority='Medium')
                if keyword_id:
                    added += 1
                else:
                    skipped += 1
            except Exception as e:
                errors.append(f"{song}: {str(e)}")
                skipped += 1
        
        return jsonify({
            'success': True,
            'artist': artist.name,
            'total_songs': result['total_songs'],
            'added': added,
            'skipped': skipped,
            'errors': errors[:5],  # First 5 errors only
            'artist_info': result['artist_info']
        })
        
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/keywords/fetch-musicbrainz/<int:artist_id>', methods=['POST'])
def fetch_musicbrainz_songs(artist_id):
    """Fetch all songs for an artist from MusicBrainz (free, no API key)"""
    try:
        # Get artist
        artist = db.get_artist(artist_id)
        if not artist:
            return jsonify({'error': 'Artist not found'}), 404
        
        # Fetch songs from MusicBrainz
        result = musicbrainz_service.get_artist_all_songs(artist.name)
        
        if 'error' in result:
            return jsonify({'error': result['error']}), 500
        
        # Add songs as keywords
        added = 0
        skipped = 0
        errors = []
        
        for song in result['songs']:
            try:
                keyword_id = db.add_keyword(song, artist_id=artist_id, auto_flag=False, priority='Medium')
                if keyword_id:
                    added += 1
                else:
                    skipped += 1
            except Exception as e:
                errors.append(f"{song}: {str(e)}")
                skipped += 1
        
        return jsonify({
            'success': True,
            'artist': artist.name,
            'total_songs': result['total_songs'],
            'added': added,
            'skipped': skipped,
            'errors': errors[:5],
            'artist_info': result['artist_info']
        })
        
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/keywords/suggest/<int:artist_id>', methods=['GET'])
def suggest_keywords(artist_id):
    """Get AI-suggested keywords based on existing video matches"""
    try:
        suggestions = keyword_learner.suggest_keywords_from_videos(artist_id=artist_id, limit=50)
        
        return jsonify({
            'success': True,
            'suggestions': suggestions,
            'count': len(suggestions)
        })
        
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/keywords/generate-patterns/<int:artist_id>', methods=['POST'])
def generate_patterns(artist_id):
    """Generate piracy pattern keywords for an artist"""
    try:
        # Get artist
        artist = db.get_artist(artist_id)
        if not artist:
            return jsonify({'error': 'Artist not found'}), 404
        
        # Generate variations
        variations = keyword_learner.suggest_artist_variations(artist.name)
        
        # Add as keywords
        added = 0
        skipped = 0
        
        for keyword in variations:
            try:
                keyword_id = db.add_keyword(keyword, artist_id=artist_id, auto_flag=False, priority='Medium')
                if keyword_id:
                    added += 1
                else:
                    skipped += 1
            except:
                skipped += 1
        
        return jsonify({
            'success': True,
            'artist': artist.name,
            'total_patterns': len(variations),
            'added': added,
            'skipped': skipped,
            'patterns': variations
        })
        
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/keywords/performance', methods=['GET'])
def keywords_performance():
    """Get performance stats for all keywords"""
    try:
        artist_id = request.args.get('artist_id', type=int)
        performances = keyword_learner.get_all_keywords_performance(artist_id=artist_id)
        
        return jsonify({
            'success': True,
            'keywords': performances,
            'count': len(performances)
        })
        
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# ============================================================================
# AUTO-UPDATE ENDPOINTS
# ============================================================================

@app.route('/api/auto-update/enable/<int:artist_id>', methods=['POST'])
def enable_auto_update(artist_id):
    """Enable automatic song updates for an artist"""
    try:
        data = request.json or {}
        frequency = data.get('frequency', 'weekly')  # daily, weekly, monthly
        source = data.get('source', 'spotify')  # spotify or musicbrainz
        
        success = auto_update_service.enable_auto_update(artist_id, frequency, source)
        
        if success:
            return jsonify({
                'success': True,
                'message': f'Auto-update enabled ({frequency} via {source})'
            })
        else:
            return jsonify({'error': 'Failed to enable auto-update'}), 500
            
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/auto-update/disable/<int:artist_id>', methods=['POST'])
def disable_auto_update(artist_id):
    """Disable automatic updates for an artist"""
    try:
        success = auto_update_service.disable_auto_update(artist_id)
        
        if success:
            return jsonify({'success': True, 'message': 'Auto-update disabled'})
        else:
            return jsonify({'error': 'Failed to disable auto-update'}), 500
            
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/auto-update/status', methods=['GET'])
def auto_update_status():
    """Get status of auto-update system"""
    try:
        status = auto_update_service.get_update_status()
        return jsonify(status)
        
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/auto-update/run/<int:artist_id>', methods=['POST'])
def run_auto_update(artist_id):
    """Manually trigger update for an artist"""
    try:
        result = auto_update_service.update_artist_songs(artist_id)
        return jsonify(result)
        
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/auto-update/run-all', methods=['POST'])
def run_all_auto_updates():
    """Run updates for all artists that need it"""
    try:
        results = auto_update_service.update_all_artists()
        
        total_new = sum(r.get('new_songs', 0) for r in results if r.get('success'))
        
        return jsonify({
            'success': True,
            'updates': results,
            'total_artists_updated': len([r for r in results if r.get('success')]),
            'total_new_songs': total_new
        })
        
    except Exception as e:
        return jsonify({'error': str(e)}), 500

if __name__ == '__main__':
    port = int(os.getenv('PORT', 5000))
    app.run(host='0.0.0.0', port=port, debug=True)
